# -*- coding: utf-8 -*-
"""
Markdown → Word 转换器（融策审计文档专用）
================================================
特性：
  - 融策配色表头（深蓝#0A1F3F底 + 铜金#C5955C字）
  - 正确处理：H1-H4标题 / 表格 / 加粗 / 引用 / 有序无序列表 / 分隔线
  - 中文字体：标题微软雅黑、正文宋体
  - 隔行浅灰底纹表格
  - A4页边距（上2.5 下2 左2.8 右2.8 cm）

用法：
  python md2docx.py "输入.md"                # 输出同名.docx
  python md2docx.py "输入.md" "输出.docx"     # 指定输出
"""
import sys
import os
import re

if sys.stdout.encoding != 'utf-8':
    try:
        sys.stdout.reconfigure(encoding='utf-8')
    except Exception:
        pass

try:
    from docx import Document
    from docx.shared import Pt, RGBColor, Cm
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement
except ImportError:
    print("缺少 python-docx，请先安装：pip install python-docx")
    sys.exit(1)

# 融策配色
NAVY = "0A1F3F"      # 深蓝
GOLD = "C5955C"      # 铜金
TEAL = "1A5C6E"      # 青绿
WARMGRAY = "F5F2EC"  # 暖灰


def set_cell_bg(cell, color_hex):
    """设置单元格背景色"""
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), color_hex)
    tcPr.append(shd)


def set_cell_font(cell, name="微软雅黑", size=10.5, bold=False, color=None):
    for p in cell.paragraphs:
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(2)
        for run in p.runs:
            run.font.name = name
            run._element.rPr.rFonts.set(qn('w:eastAsia'), name)
            run.font.size = Pt(size)
            run.font.bold = bold
            if color:
                run.font.color.rgb = RGBColor.from_string(color)


def add_runs_with_bold(paragraph, text):
    """处理 **加粗** 标记，也去掉行内 `代码` 反引号"""
    text = text.replace('`', '')
    parts = re.split(r'(\*\*[^*]+\*\*)', text)
    for part in parts:
        if part.startswith('**') and part.endswith('**') and len(part) > 4:
            run = paragraph.add_run(part[2:-2])
            run.bold = True
        elif part:
            paragraph.add_run(part)


def setup_document():
    doc = Document()
    # 正文样式：宋体小四
    normal = doc.styles['Normal']
    normal.font.name = '宋体'
    normal._element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')
    normal.font.size = Pt(12)
    # 页边距
    sec = doc.sections[0]
    sec.top_margin = Cm(2.5)
    sec.bottom_margin = Cm(2)
    sec.left_margin = Cm(2.8)
    sec.right_margin = Cm(2.8)
    return doc


def style_heading(paragraph, size, color=NAVY):
    for run in paragraph.runs:
        run.font.name = '微软雅黑'
        run._element.rPr.rFonts.set(qn('w:eastAsia'), '微软雅黑')
        run.font.size = Pt(size)
        run.font.color.rgb = RGBColor.from_string(color)
        run.font.bold = True


def parse_table_block(lines, start):
    """从start开始收集表格行，返回(rows, next_index)"""
    table_lines = []
    i = start
    while i < len(lines) and '|' in lines[i]:
        table_lines.append(lines[i].strip())
        i += 1
    # 过滤分隔行 |---|---|
    data_lines = []
    for tl in table_lines:
        stripped = tl.replace('|', '').replace('-', '').replace(':', '').strip()
        if stripped == '':
            continue  # 分隔行
        data_lines.append(tl)
    rows = []
    for tl in data_lines:
        cells = [c.strip() for c in tl.split('|')]
        if cells and cells[0] == '':
            cells = cells[1:]
        if cells and cells[-1] == '':
            cells = cells[:-1]
        if cells:
            rows.append(cells)
    return rows, i


def md_to_docx(md_path, out_path):
    with open(md_path, 'r', encoding='utf-8') as f:
        lines = f.readlines()

    doc = setup_document()
    i = 0
    n = len(lines)
    while i < n:
        raw = lines[i].rstrip('\n')
        line = raw.rstrip()

        if not line.strip():
            i += 1
            continue

        # 标题
        if line.startswith('#### '):
            p = doc.add_paragraph()
            add_runs_with_bold(p, line[5:])
            style_heading(p, 12, TEAL)
            i += 1
            continue
        if line.startswith('### '):
            p = doc.add_paragraph()
            add_runs_with_bold(p, line[4:])
            style_heading(p, 13, TEAL)
            i += 1
            continue
        if line.startswith('## '):
            p = doc.add_paragraph()
            add_runs_with_bold(p, line[3:])
            style_heading(p, 15, NAVY)
            i += 1
            continue
        if line.startswith('# '):
            p = doc.add_paragraph()
            add_runs_with_bold(p, line[2:])
            style_heading(p, 20, NAVY)
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            i += 1
            continue

        # 分隔线
        if re.match(r'^-{3,}$', line.strip()) or re.match(r'^\*{3,}$', line.strip()):
            p = doc.add_paragraph()
            run = p.add_run('─' * 40)
            run.font.color.rgb = RGBColor.from_string(GOLD)
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            i += 1
            continue

        # 代码块 ```
        if line.strip().startswith('```'):
            i += 1
            code_lines = []
            while i < n and not lines[i].strip().startswith('```'):
                code_lines.append(lines[i].rstrip('\n'))
                i += 1
            i += 1  # 跳过结尾```
            for cl in code_lines:
                p = doc.add_paragraph()
                run = p.add_run(cl if cl else ' ')
                run.font.name = 'Consolas'
                run.font.size = Pt(10)
                p.paragraph_format.left_indent = Cm(0.5)
            continue

        # 引用
        if line.startswith('> '):
            p = doc.add_paragraph()
            add_runs_with_bold(p, line[2:])
            p.paragraph_format.left_indent = Cm(1)
            for run in p.runs:
                run.font.color.rgb = RGBColor(0x66, 0x66, 0x66)
                run.font.italic = True
            i += 1
            continue

        # 表格
        if '|' in line:
            rows, next_i = parse_table_block(lines, i)
            if rows:
                max_cols = max(len(r) for r in rows)
                table = doc.add_table(rows=len(rows), cols=max_cols)
                table.alignment = WD_TABLE_ALIGNMENT.CENTER
                table.style = 'Table Grid'
                for r_idx, row in enumerate(rows):
                    for c_idx in range(max_cols):
                        cell = table.rows[r_idx].cells[c_idx]
                        text = row[c_idx] if c_idx < len(row) else ''
                        # 清理加粗标记和反引号
                        text = text.replace('**', '').replace('`', '')
                        cell.text = text
                        if r_idx == 0:
                            set_cell_bg(cell, NAVY)
                            set_cell_font(cell, "微软雅黑", 10.5, bold=True, color=GOLD)
                        else:
                            if r_idx % 2 == 0:
                                set_cell_bg(cell, WARMGRAY)
                            set_cell_font(cell, "微软雅黑", 10, bold=False)
                doc.add_paragraph()  # 表后空行
            i = next_i
            continue

        # 无序列表（含缩进）
        m_ul = re.match(r'^(\s*)[-*]\s+(.+)', raw)
        if m_ul:
            text = m_ul.group(2)
            # 处理 [ ] / [x] 复选框
            text = re.sub(r'^\[[ xX]\]\s*', '☐ ', text)
            p = doc.add_paragraph(style='List Bullet')
            add_runs_with_bold(p, text)
            i += 1
            continue

        # 有序列表
        m_ol = re.match(r'^(\s*)(\d+)\.\s+(.+)', raw)
        if m_ol:
            p = doc.add_paragraph(style='List Number')
            add_runs_with_bold(p, m_ol.group(3))
            i += 1
            continue

        # 普通段落
        p = doc.add_paragraph()
        add_runs_with_bold(p, line)
        i += 1

    doc.save(out_path)
    return out_path


def main():
    if len(sys.argv) < 2:
        print("用法：python md2docx.py \"输入.md\" [\"输出.docx\"]")
        sys.exit(1)
    md_path = sys.argv[1]
    if not os.path.exists(md_path):
        print(f"❌ 找不到文件：{md_path}")
        sys.exit(1)
    if len(sys.argv) >= 3:
        out_path = sys.argv[2]
    else:
        out_path = os.path.splitext(md_path)[0] + '.docx'
    md_to_docx(md_path, out_path)
    size_kb = os.path.getsize(out_path) / 1024
    print(f"✅ 已转换：{out_path}  ({size_kb:.0f} KB)")


if __name__ == "__main__":
    main()
